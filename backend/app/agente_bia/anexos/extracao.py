"""Saneamento de imagem e extração de texto. Roda no UPLOAD, nunca no envio.

- Imagem: decodificada e RECODIFICADA pelo Pillow no mesmo formato. Isso
  descarta EXIF (GPS, câmera, data), ICC e qualquer byte além da imagem, e
  recusa arquivo que não decodifica de verdade. Limite de pixels contra bomba
  de descompressão. A orientação EXIF é aplicada antes de descartá-la.
- PDF: texto por página com `pypdf`. Sem texto extraível (digitalizado) é
  RECUSADO: a V1 não faz OCR, e a Bia nunca recebe um PDF que não consegue ler.
- DOCX: parágrafos e tabelas com `python-docx`.
- TXT/CSV: UTF-8, com recaída para cp1252; controle removido.

Nenhum conteúdo extraído vai para log.
"""
from __future__ import annotations

import csv
import io
import warnings
from dataclasses import dataclass
from typing import Optional

from app.agente_bia.anexos.dominio import (
    MAX_CARACTERES_POR_DOCUMENTO,
    MAX_LADO_ENVIO_IMAGEM,
    MAX_LADO_IMAGEM,
    MAX_PAGINAS_PDF,
    MAX_PIXELS_IMAGEM,
    AnexoRecusado,
    TipoAceito,
)


@dataclass(frozen=True)
class Processado:
    dados: bytes                 # bytes que ficam guardados (imagem saneada ou original)
    texto: Optional[str]         # texto extraído (documentos)
    paginas: Optional[int]
    largura: Optional[int]
    altura: Optional[int]
    truncado: bool


def _imagem(dados: bytes, tipo: TipoAceito) -> Processado:
    from PIL import Image, ImageOps

    try:
        # O aviso de bomba do Pillow vira erro AQUI (sem mexer no global do processo),
        # e o limite de pixels é conferido no cabeçalho, ANTES de decodificar:
        # `Image.MAX_IMAGE_PIXELS` sozinho só avisa até 2x o limite.
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(dados)) as bruta:
                if bruta.width * bruta.height > MAX_PIXELS_IMAGEM:
                    raise AnexoRecusado("imagem_grande", "Imagem com resolução acima do limite.")
                if bruta.width > MAX_LADO_IMAGEM or bruta.height > MAX_LADO_IMAGEM:
                    raise AnexoRecusado("imagem_grande", f"Imagem acima de {MAX_LADO_IMAGEM} px de lado.")
                bruta.verify()
            with Image.open(io.BytesIO(dados)) as original:
                img = _saneada(original, ImageOps)
            formato = {"png": "PNG", "jpeg": "JPEG", "webp": "WEBP"}[tipo.chave]
            if formato == "JPEG" and img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            saida = io.BytesIO()
            kwargs = {"quality": 90} if formato in ("JPEG", "WEBP") else {"optimize": True}
            img.save(saida, format=formato, **kwargs)  # sem exif=, sem icc_profile=
            return Processado(saida.getvalue(), None, None, img.width, img.height, False)
    except AnexoRecusado:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise AnexoRecusado("imagem_grande", "Imagem com resolução acima do limite.") from exc
    except Exception as exc:  # noqa: BLE001
        raise AnexoRecusado("imagem_invalida", "A imagem não pôde ser lida.") from exc


def _saneada(img, ImageOps):
    """Orientação aplicada e lado maior reduzido para o que a visão usa de fato."""
    img = ImageOps.exif_transpose(img)
    if max(img.width, img.height) > MAX_LADO_ENVIO_IMAGEM:
        img = img.copy()
        img.thumbnail((MAX_LADO_ENVIO_IMAGEM, MAX_LADO_ENVIO_IMAGEM))
    return img


def _cortar(texto: str) -> tuple[str, bool]:
    texto = "".join(ch for ch in texto if ch in "\n\t" or ord(ch) >= 32).strip()
    if len(texto) > MAX_CARACTERES_POR_DOCUMENTO:
        return texto[:MAX_CARACTERES_POR_DOCUMENTO], True
    return texto, False


def _pdf(dados: bytes) -> Processado:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover — ambiente sem pypdf
        raise AnexoRecusado("pdf_indisponivel", "Leitura de PDF indisponível neste servidor.") from exc
    try:
        leitor = PdfReader(io.BytesIO(dados))
        if leitor.is_encrypted:
            raise AnexoRecusado("pdf_protegido", "O PDF está protegido por senha.")
        paginas = len(leitor.pages)
        if paginas > MAX_PAGINAS_PDF:
            raise AnexoRecusado("pdf_grande", f"PDF acima de {MAX_PAGINAS_PDF} páginas.")
        partes, total = [], 0
        for i, pagina in enumerate(leitor.pages):
            parte = f"[página {i + 1}]\n{(pagina.extract_text() or '').strip()}"
            partes.append(parte)
            total += len(parte)
            if total > MAX_CARACTERES_POR_DOCUMENTO:
                break  # o resto seria cortado de qualquer forma
    except AnexoRecusado:
        raise
    except Exception as exc:  # noqa: BLE001
        raise AnexoRecusado("pdf_invalido", "O PDF não pôde ser lido.") from exc
    texto, truncado = _cortar("\n\n".join(partes))
    if len(texto.replace("[página", "").strip()) < 20 or not any(
            p.split("\n", 1)[-1].strip() for p in partes):
        raise AnexoRecusado("pdf_sem_texto",
                            "O PDF não tem texto extraível (parece digitalizado). Envie o texto ou uma imagem da página.")
    return Processado(dados, texto, paginas, None, None, truncado)


def _docx(dados: bytes) -> Processado:
    try:
        import docx  # python-docx
        doc = docx.Document(io.BytesIO(dados))
        linhas = [p.text for p in doc.paragraphs if p.text.strip()]
        for tabela in doc.tables:
            for linha in tabela.rows:
                linhas.append(" | ".join(c.text.strip() for c in linha.cells))
    except Exception as exc:  # noqa: BLE001
        raise AnexoRecusado("docx_invalido", "O DOCX não pôde ser lido.") from exc
    texto, truncado = _cortar("\n".join(linhas))
    if not texto:
        raise AnexoRecusado("documento_vazio", "O DOCX não tem texto.")
    return Processado(dados, texto, None, None, None, truncado)


def _decodificar(dados: bytes) -> str:
    try:
        return dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        return dados.decode("cp1252", errors="replace")


def _texto(dados: bytes, tipo: TipoAceito) -> Processado:
    texto = _decodificar(dados)
    if tipo.chave == "csv":
        try:
            linhas = list(csv.reader(io.StringIO(texto)))
        except csv.Error as exc:
            raise AnexoRecusado("csv_invalido", "O CSV não pôde ser lido.") from exc
        texto = "\n".join(" | ".join(c.strip() for c in l) for l in linhas)
    texto, truncado = _cortar(texto)
    if not texto:
        raise AnexoRecusado("documento_vazio", "O arquivo não tem texto.")
    return Processado(dados, texto, None, None, None, truncado)


def processar(dados: bytes, tipo: TipoAceito) -> Processado:
    if tipo.familia == "imagem":
        return _imagem(dados, tipo)
    if tipo.chave == "pdf":
        return _pdf(dados)
    if tipo.chave == "docx":
        return _docx(dados)
    return _texto(dados, tipo)
