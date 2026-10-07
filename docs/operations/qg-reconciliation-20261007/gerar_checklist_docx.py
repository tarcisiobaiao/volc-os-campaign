#!/usr/bin/env python3
"""Gera o checklist real de fechamento (JSON + DOCX) a partir do Roadmap Vivo.

O texto de cada item vem de `checklist_itens.py`; o status vem do
`volc-os-workbook/ROADMAP-VIVO.json`. O gerador recusa rodar se os dois
divergirem: item com ID inexistente, item já concluído ou reservado, tarefa
aberta das ondas A–C sem item, ou ID repetido.

Uso (com o Python que tem python-docx, o mesmo do volc-os-workbook/build.py):

    backend/.venv/bin/python docs/operations/qg-reconciliation-20261007/gerar_checklist_docx.py
"""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(HERE))
from checklist_itens import FAIXAS, ITENS  # noqa: E402

ROADMAP = ROOT / "volc-os-workbook" / "ROADMAP-VIVO.json"
SAIDA = ROOT / "entregaveis"
NOME = "VOLC-OS-CHECKLIST-REAL-2026-10-07"
DATA = date(2026, 10, 7)
ABERTOS = {"partial", "todo", "risk"}
FORA_DO_SPRINT = "D — "  # prefixo das ondas que o próprio roadmap estaciona

# Mesmo padrão visual do volc-os-workbook/build.py.
NAVY, DEEP, ORANGE, INK, MUTED, LINE = "0B1020", "0D47A1", "FF4D20", "1A1C1E", "667085", "DDE4EE"
STATUS = {"partial": "Parcial", "todo": "A fazer", "risk": "Com risco"}

DECISOES = [
    ("Variante v12_04",
     "A sprint last-mile trouxe outra implementação da RPC v12_04 (48d9fc0, 79770a5, c299274), que aponta defeitos na projeção legada. Ficou fora; compare com a versão adjudicada antes da migration oficial."),
    ("Autoridade de nascimento Hermes",
     "c689cdb é outra implementação de P09-T17 e nunca entrou na linha v2. Decida se substitui, complementa ou sai."),
    ("Attention Engine",
     "Está só no worktree principal, sem commit. A sessão de origem precisa fechar o trabalho numa branch própria antes de qualquer integração."),
    ("Âncoras de CTA",
     "Ficou a regra H1+slug da linha v2, que alimenta o portão do destino pago. Ela produz textos como \"Ver o guia de social\"; o Editorial V2 preferia o H1 inteiro."),
    ("Exportação DOCX do QG",
     "O botão DOCX do QG entrega um arquivo fixo de entregaveis/, não o roadmap vivo. Este checklist foi gerado por outro caminho."),
    ("Promoção ao main",
     "Nada foi mesclado no main nem enviado ao remoto. É o primeiro item do checklist (P01-T12)."),
]


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, check=True,
                          capture_output=True, text=True).stdout.strip()


def carregar() -> tuple[dict, list[dict], list[dict]]:
    bruto = ROADMAP.read_bytes()
    doc = json.loads(bruto)
    tarefas = {}
    for ini in doc["initiatives"]:
        for t in ini["tasks"]:
            tarefas[t["id"]] = {**t, "_onda": ini.get("wave", ""), "_iniciativa": ini["id"]}

    ids = [i[0] for i in ITENS]
    erros = []
    if len(ids) != len(set(ids)):
        erros.append("ID repetido no checklist")
    for tid in ids:
        if tid not in tarefas:
            erros.append(f"{tid} não existe no roadmap")
        elif tarefas[tid]["status"] not in ABERTOS:
            erros.append(f"{tid} está {tarefas[tid]['status']} no roadmap")
    abertas = {k for k, t in tarefas.items()
               if t["status"] in ABERTOS and not t["_onda"].startswith(FORA_DO_SPRINT)}
    faltando = abertas - set(ids)
    sobrando = set(ids) - abertas
    if faltando:
        erros.append(f"tarefas abertas sem item: {sorted(faltando)}")
    if sobrando:
        erros.append(f"itens fora das ondas A–C: {sorted(sobrando)}")
    for tid, faixa, *_ in ITENS:
        if faixa not in FAIXAS:
            erros.append(f"{tid} com faixa {faixa}")
    if erros:
        raise SystemExit("Checklist diverge do roadmap:\n- " + "\n- ".join(erros))

    itens = []
    for ordem, (tid, faixa, titulo, falta, concluido, dep) in enumerate(ITENS, start=1):
        itens.append({
            "ordem": ordem, "id": tid, "status": tarefas[tid]["status"],
            "faixa": faixa, "faixa_nome": FAIXAS[faixa], "titulo": titulo,
            "falta": falta, "concluido_quando": concluido,
            "dependencia": dep or None,
        })
    fora = [{"id": k, "status": t["status"], "titulo": t["title"], "onda": t["_onda"],
             "motivo": "Onda D: o roadmap estaciona esta frente até haver dados, ledger e modelo para ela."}
            for k, t in tarefas.items()
            if t["status"] in ABERTOS and t["_onda"].startswith(FORA_DO_SPRINT)]
    meta = {
        "titulo": "VOLC-OS — Checklist real de fechamento",
        "data": DATA.isoformat(),
        "branch": git("rev-parse", "--abbrev-ref", "HEAD"),
        "head": git("rev-parse", "--short", "HEAD"),
        "roadmap_commit": git("log", "-1", "--format=%h", "--", str(ROADMAP.relative_to(ROOT))),
        "roadmap_sha256": hashlib.sha256(bruto).hexdigest(),
        "itens_no_sprint": len(itens),
        "fora_do_sprint": len(fora),
    }
    return meta, itens, fora


def estilo(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.1
    for nome, tam, cor in (("Title", 22, NAVY), ("Heading 1", 14, DEEP), ("Heading 2", 11, ORANGE)):
        s = doc.styles[nome]
        s.font.name = "Space Grotesk"
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(cor)
        s.paragraph_format.space_before = Pt(12 if nome != "Title" else 0)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)


def rodape(doc: Document, meta: dict) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"VOLC O.S.  •  checklist {meta['data']}  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
    txt = OxmlElement("w:instrText"); txt.set(qn("xml:space"), "preserve"); txt.text = "PAGE"
    fim = OxmlElement("w:fldChar"); fim.set(qn("w:fldCharType"), "end")
    run._r.extend((ini, txt, fim))


def linha(doc: Document, rotulo: str, texto: str, cor: str = MUTED) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    r = p.add_run(rotulo)
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(cor)
    r2 = p.add_run(texto)
    r2.font.size = Pt(9.5)


def gerar_docx(meta: dict, itens: list[dict], fora: list[dict], destino: Path) -> None:
    doc = Document()
    estilo(doc)
    rodape(doc, meta)
    doc.add_paragraph(meta["titulo"], style="Title")
    p = doc.add_paragraph()
    r = p.add_run(f"07/10/2026  •  branch {meta['branch']}  •  HEAD {meta['head']}  •  roadmap {meta['roadmap_commit']}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph(
        "A lista sai do produto integrado nessa branch e do Roadmap Vivo reconciliado no mesmo dia: "
        f"são as {meta['itens_no_sprint']} tarefas ainda abertas das ondas A, B e C, na ordem em que fazem sentido.")

    faixa_atual = None
    for item in itens:
        if item["faixa"] != faixa_atual:
            faixa_atual = item["faixa"]
            doc.add_paragraph(f"{faixa_atual}. {item['faixa_nome']}", style="Heading 1")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.keep_with_next = True
        caixa = p.add_run("☐  ")
        caixa.font.size = Pt(12)
        rid = p.add_run(f"{item['id']}  ")
        rid.bold = True
        rid.font.color.rgb = RGBColor.from_string(DEEP)
        rt = p.add_run(item["titulo"])
        rt.bold = True
        rs = p.add_run(f"  ·  {STATUS[item['status']]}")
        rs.font.size = Pt(8.5)
        rs.font.color.rgb = RGBColor.from_string(MUTED)
        d = doc.add_paragraph(item["falta"])
        d.paragraph_format.left_indent = Cm(0.75)
        d.paragraph_format.keep_with_next = True
        linha(doc, "Concluído quando: ", item["concluido_quando"], DEEP)
        if item["dependencia"]:
            linha(doc, "Depende de: ", item["dependencia"], ORANGE)

    doc.add_page_break()
    doc.add_paragraph("Riscos e decisões humanas", style="Heading 1")
    for titulo, texto in DECISOES:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        a = p.add_run(f"{titulo}. ")
        a.bold = True
        p.add_run(texto)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    a = p.add_run("Fora deste sprint. ")
    a.bold = True
    p.add_run("A onda D do roadmap (preditivo, Meridian e Gold Site Factory) fica estacionada: "
              + ", ".join(f["id"] for f in fora) + ".")
    doc.save(destino)


def main() -> None:
    meta, itens, fora = carregar()
    SAIDA.mkdir(exist_ok=True)
    js = SAIDA / f"{NOME}.json"
    js.write_text(json.dumps({"meta": meta, "itens": itens, "fora_do_sprint": fora,
                              "decisoes_humanas": [{"titulo": t, "texto": x} for t, x in DECISOES]},
                             ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    gerar_docx(meta, itens, fora, SAIDA / f"{NOME}.docx")
    print(json.dumps({"json": str(js), "docx": str(SAIDA / f"{NOME}.docx"),
                      "itens": len(itens), "fora_do_sprint": len(fora)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
